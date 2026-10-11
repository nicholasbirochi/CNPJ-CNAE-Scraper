"""Lê CNPJs de data/ e gera output.xlsx usando a BrasilAPI gratuita."""

import argparse
import csv
from datetime import datetime, timezone, timedelta
import json
from pathlib import Path
import re
import sys
import ssl
import sqlite3
import tempfile
import time
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

import certifi
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from consultar import ErroConsulta, normalizar_cnpj

RAIZ = Path(__file__).resolve().parent
CABECALHO = ["CNPJ", "CNAE principal", "CNAEs secundários"]
PADRAO_CNPJ = re.compile(
    r"(?<![A-Z0-9])(?:[A-Z0-9]{2}\.[A-Z0-9]{3}\.[A-Z0-9]{3}/"
    r"[A-Z0-9]{4}-[0-9]{2}|[A-Z0-9]{8}/[A-Z0-9]{4}-[0-9]{2}|[A-Z0-9]{12}[0-9]{2})(?![A-Z0-9])", re.I
)


class LimiteAPI(ErroConsulta):
    """Interrompe o lote quando o provedor bloqueia ou limita o acesso."""


def _celula(valor) -> str:
    if isinstance(valor, (float, int)) and not isinstance(valor, bool):
        if valor == int(valor) and 0 <= valor < 10**14:
            return str(int(valor)).zfill(14)
    return str(valor).strip() if valor is not None else ""


def _coluna_cnpj(linhas: list) -> list[str]:
    for n, linha in enumerate(linhas[:20]):
        colunas = [i for i, v in enumerate(linha) if "CNPJ" in str(v).upper()]
        if len(colunas) == 1:
            coluna = colunas[0]
            return [_celula(row[coluna]) for row in linhas[n + 1:]
                    if len(row) > coluna and _celula(row[coluna])]
    encontrados = []
    for linha in linhas:
        for valor in linha:
            texto = _celula(valor)
            encontrados.extend(PADRAO_CNPJ.findall(texto))
    return encontrados


def ler_cnpjs(entrada: Path, saida: Path) -> list[str]:
    if not entrada.exists():
        raise ErroConsulta(f"Pasta ou arquivo de entrada não encontrado: {entrada}")
    arquivos = sorted(entrada.rglob("*")) if entrada.is_dir() else [entrada]
    cnpjs = []
    for arquivo in arquivos:
        if not arquivo.is_file() or arquivo.resolve() == saida.resolve() or arquivo.name.startswith("~$"):
            continue
        if arquivo.suffix.lower() == ".xlsx":
            livro = load_workbook(arquivo, read_only=True, data_only=True)
            try:
                for folha in livro:
                    cnpjs.extend(_coluna_cnpj(list(folha.iter_rows(values_only=True))))
            finally:
                livro.close()
        elif arquivo.suffix.lower() in {".csv", ".txt", ".tsv"}:
            bruto = arquivo.read_bytes()
            try:
                texto = bruto.decode("utf-8-sig")
            except UnicodeDecodeError:
                texto = bruto.decode("cp1252")
            if arquivo.suffix.lower() == ".txt":
                linhas = [[linha] for linha in texto.splitlines()]
            else:
                try:
                    dialeto = csv.Sniffer().sniff(texto[:8192], delimiters=",;\t")
                except csv.Error:
                    dialeto = csv.excel
                linhas = list(csv.reader(texto.splitlines(), dialect=dialeto))
            cnpjs.extend(_coluna_cnpj(linhas))
        elif arquivo.suffix.lower() == ".xls":
            raise ErroConsulta(f"Converta {arquivo.name} de .xls para .xlsx antes de consultar.")
    if not cnpjs:
        raise ErroConsulta("Nenhum CNPJ encontrado. Use XLSX, CSV, TSV ou TXT em data/.")
    return cnpjs


def _atividade(codigo, descricao) -> dict:
    texto = str(codigo)
    if not re.fullmatch(r"[0-9]{1,7}", texto) or int(texto) == 0:
        raise ErroConsulta("BrasilAPI retornou um código CNAE inválido.")
    if not isinstance(descricao, str) or not descricao.strip():
        raise ErroConsulta("BrasilAPI retornou CNAE sem descrição.")
    texto = texto.zfill(7)
    return {"codigo": f"{texto[:2]}.{texto[2:4]}-{texto[4]}-{texto[5:]}",
            "descricao": descricao.strip()}


def extrair_api(dados: dict, cnpj: str) -> dict:
    if not isinstance(dados, dict) or normalizar_cnpj(str(dados.get("cnpj", ""))) != cnpj:
        raise ErroConsulta("BrasilAPI retornou dados de outro CNPJ ou resposta inválida.")
    if not isinstance(dados.get("cnaes_secundarios"), list):
        raise ErroConsulta("BrasilAPI não retornou o campo de CNAEs secundários.")
    principal = _atividade(dados.get("cnae_fiscal"), dados.get("cnae_fiscal_descricao"))
    secundarios = []
    for item in dados["cnaes_secundarios"]:
        if not isinstance(item, dict):
            raise ErroConsulta("BrasilAPI retornou um CNAE secundário inválido.")
        if item.get("codigo") in (0, "0") and not item.get("descricao"):
            continue
        atividade = _atividade(item.get("codigo"), item.get("descricao"))
        if atividade not in secundarios:
            secundarios.append(atividade)
    return {"cnae_principal": principal, "cnaes_secundarios": secundarios}


def consultar_api(cnpj: str) -> dict:
    pedido = Request(f"https://brasilapi.com.br/api/cnpj/v1/{cnpj}", headers={
        "Accept": "application/json", "User-Agent": "CNPJ-CNAE-Scraper/1.0"
    })
    try:
        with urlopen(pedido, timeout=30, context=ssl.create_default_context(cafile=certifi.where())) as resposta:
            return extrair_api(json.load(resposta), cnpj)
    except HTTPError as erro:
        if erro.code in (403, 429):
            raise LimiteAPI(f"BrasilAPI limitou o acesso (HTTP {erro.code}). Tente mais tarde.") from erro
        if erro.code == 404:
            raise ErroConsulta("CNPJ não encontrado na base da BrasilAPI.") from erro
        raise ErroConsulta(f"BrasilAPI indisponível (HTTP {erro.code}).") from erro
    except (URLError, TimeoutError, json.JSONDecodeError) as erro:
        raise ErroConsulta("Falha de conexão ou resposta inválida da BrasilAPI.") from erro


def _resumo(atividade: dict) -> str:
    return f"{atividade['codigo']} - {atividade['descricao']}"


def montar_linhas(cnpjs: list[str], resultados: dict, erros: dict) -> list[list[str]]:
    linhas = []
    for original in cnpjs:
        try:
            cnpj = normalizar_cnpj(original)
        except ErroConsulta as erro:
            linhas.append([original, f"ERRO: {erro}", "Não consultado"])
            continue
        formatado = f"{cnpj[:2]}.{cnpj[2:5]}.{cnpj[5:8]}/{cnpj[8:12]}-{cnpj[12:]}"
        if cnpj in resultados:
            resultado = resultados[cnpj]
            secundarios = "\n".join(_resumo(a) for a in resultado["cnaes_secundarios"])
            linhas.append([formatado, _resumo(resultado["cnae_principal"]), secundarios or "Não informada"])
        else:
            linhas.append([formatado, f"ERRO: {erros[cnpj]}" if cnpj in erros else "PENDENTE", "Não consultado"])
    return linhas


def salvar_excel(caminho: Path, linhas: list[list[str]]) -> None:
    # O gerador portátil usa openpyxl; o arquivo é substituído somente após salvar com sucesso.
    livro = Workbook()
    folha = livro.active
    folha.title = "CNAEs"
    folha.sheet_view.showGridLines = False
    folha.append(CABECALHO)
    for linha in linhas:
        folha.append(linha)
    for celula in folha[1]:
        celula.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="17365D")
        celula.alignment = Alignment(horizontal="center", vertical="center")
    folha.row_dimensions[1].height = 28
    for row in folha.iter_rows(min_row=2):
        for celula in row:
            celula.data_type = "s"  # Dados externos permanecem texto, inclusive quando começam por '='.
            celula.number_format = "@"
            celula.font = Font(name="Arial", size=11)
            celula.alignment = Alignment(wrap_text=True, vertical="top")
        linhas_visuais = max(sum(max(1, (len(parte) + 64) // 65)
                                 for parte in str(c.value).split("\n")) for c in row)
        folha.row_dimensions[row[0].row].height = min(409, 16 * linhas_visuais + 8)
    for coluna, largura in {"A": 24, "B": 72, "C": 95}.items():
        folha.column_dimensions[coluna].width = largura
    folha.freeze_panes = "B2"
    folha.auto_filter.ref = folha.dimensions
    livro.properties.description = ("Fonte: Receita Federal, via BrasilAPI ou Dados Abertos do CNPJ. "
                                    + datetime.now(timezone.utc).isoformat())
    caminho.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=caminho.parent, suffix=".xlsx", delete=False) as arquivo:
        temporario = Path(arquivo.name)
    try:
        livro.save(temporario)
        temporario.replace(caminho)
    finally:
        temporario.unlink(missing_ok=True)


def processar(entrada: Path, saida: Path, intervalo: float = 2, atualizar: bool = False, somente_cache: bool = False) -> int:
    cnpjs = ler_cnpjs(entrada, saida)
    if saida.exists():
        existente = load_workbook(saida, read_only=True)
        try:
            if list(next(existente.active.iter_rows(max_row=1, values_only=True))) != CABECALHO:
                raise ErroConsulta("A saída já existe e não é deste projeto. Escolha outro nome.")
        finally:
            existente.close()
    saida.parent.mkdir(parents=True, exist_ok=True)
    cache = sqlite3.connect(saida.with_suffix(".cache.sqlite3"))
    cache.execute("CREATE TABLE IF NOT EXISTS cnaes (cnpj TEXT PRIMARY KEY, consultado_em REAL, resultado TEXT)")
    agora = time.time()
    resultados, erros = {}, {}
    validos = []
    vistos = set()
    invalidos = 0
    for original in cnpjs:
        try:
            cnpj = normalizar_cnpj(original)
        except ErroConsulta:
            invalidos += 1
            continue
        if cnpj not in vistos:
            validos.append(cnpj)
            vistos.add(cnpj)
    if not atualizar:
        for cnpj, resultado in cache.execute("SELECT cnpj, resultado FROM cnaes WHERE consultado_em > ?", (agora - 30 * 86400,)):
            if cnpj in vistos:
                resultados[cnpj] = json.loads(resultado)
    salvar_excel(saida, montar_linhas(cnpjs, resultados, erros))
    consultas = [cnpj for cnpj in validos if cnpj not in resultados]
    print(f"{len(cnpjs)} linhas, {len(validos)} CNPJs únicos válidos, {len(consultas)} consultas necessárias.", flush=True)
    if somente_cache:
        cache.close()
        return 1 if consultas or invalidos else 0
    interrompido = False
    try:
        for i, cnpj in enumerate(consultas):
            if i:
                time.sleep(intervalo)
            try:
                resultados[cnpj] = consultar_api(cnpj)
                cache.execute("INSERT OR REPLACE INTO cnaes VALUES (?, ?, ?)",
                              (cnpj, time.time(), json.dumps(resultados[cnpj], ensure_ascii=False)))
                cache.commit()
            except LimiteAPI as erro:
                erros[cnpj] = str(erro)
                print(str(erro), file=sys.stderr, flush=True)
                interrompido = True
                break
            except ErroConsulta as erro:
                erros[cnpj] = str(erro)
            print(f"Consulta {i + 1}/{len(consultas)}: {'erro' if cnpj in erros else 'concluída'}", flush=True)
            if (i + 1) % 300 == 0:
                salvar_excel(saida, montar_linhas(cnpjs, resultados, erros))
    except KeyboardInterrupt:
        interrompido = True
        print("Interrompido. Progresso salvo; execute novamente para continuar.", file=sys.stderr)
    finally:
        cache.close()
        salvar_excel(saida, montar_linhas(cnpjs, resultados, erros))
    pendentes = len(validos) - len(resultados)
    print(f"Saída: {saida}. {len(resultados)} CNPJs preenchidos; {pendentes} com erro ou pendentes.")
    return 1 if interrompido or erros or invalidos else 0


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("entrada", nargs="?", type=Path, default=RAIZ / "data")
    parser.add_argument("--saida", type=Path, default=RAIZ / "output.xlsx")
    parser.add_argument("--intervalo", type=float, default=2, help="Segundos entre consultas (mínimo 1)")
    parser.add_argument("--atualizar", action="store_true", help="Ignorar cache de consultas anteriores")
    parser.add_argument("--exportar", action="store_true", help="Atualizar somente o Excel a partir do cache, sem consultas")
    args = parser.parse_args(argv)
    if args.intervalo < 1:
        parser.error("--intervalo deve ser pelo menos 1 segundo")
    try:
        return processar(args.entrada, args.saida, args.intervalo, args.atualizar, args.exportar)
    except (ErroConsulta, OSError, ValueError) as erro:
        print(f"Erro: {erro}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
