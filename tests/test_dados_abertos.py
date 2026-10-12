from contextlib import redirect_stderr, redirect_stdout
import io
import json
from pathlib import Path
import sqlite3
import tempfile
import time
import unittest
from unittest.mock import patch
import zipfile

from openpyxl import load_workbook

from consultar import ErroConsulta
from dados_abertos import _listar, baixar, buscar, importar, ler_cnaes, montar_resultado

CNAES = ('"6201501";"Desenvolvimento  de programas"\n"6204000";"Consultoria em tecnologia da informação"\n'
         '"0111301";"Cultivo de arroz"\n')


def linha(cnpj: str, principal: str, secundarios: str) -> str:
    campos = [cnpj[:8], cnpj[8:12], cnpj[12:], "1", "", "02", "20200101", "00", "", "", "20200101",
              principal, secundarios] + [""] * 17
    return ";".join(f'"{campo}"' for campo in campos)


def criar_zips(pasta: Path, linhas: list[str]) -> None:
    with zipfile.ZipFile(pasta / "Cnaes.zip", "w", zipfile.ZIP_DEFLATED) as pacote:
        pacote.writestr("F.K03200$Z.D60912.CNAECSV", CNAES.encode("latin-1"))
    with zipfile.ZipFile(pasta / "Estabelecimentos0.zip", "w", zipfile.ZIP_DEFLATED) as pacote:
        pacote.writestr("K3241.K03200Y0.D60912.ESTABELE", "\n".join(linhas).encode("latin-1"))


class TestArquivos(unittest.TestCase):
    def test_tabela_de_cnaes_em_latin1_normaliza_espacos(self):
        with tempfile.TemporaryDirectory() as pasta:
            criar_zips(Path(pasta), [])
            tabela = ler_cnaes(Path(pasta) / "Cnaes.zip")
        self.assertEqual(tabela["6201501"], "Desenvolvimento de programas")
        self.assertEqual(tabela["6204000"], "Consultoria em tecnologia da informação")

    def test_busca_apenas_os_cnpjs_procurados(self):
        with tempfile.TemporaryDirectory() as pasta:
            criar_zips(Path(pasta), [linha("11222333000181", "6201501", ""),
                                     linha("00000000000191", "6204000", "6201501"),
                                     linha("12ABC34501DE35", "111301", "")])
            encontrados = buscar(Path(pasta) / "Estabelecimentos0.zip", {b"00000000000191", b"12ABC34501DE35"})
        self.assertEqual(set(encontrados), {"00000000000191", "12ABC34501DE35"})
        self.assertEqual(encontrados["00000000000191"][11:13], ["6204000", "6201501"])

    def test_layout_desconhecido_e_erro(self):
        with tempfile.TemporaryDirectory() as pasta:
            criar_zips(Path(pasta), ["CNPJ;CNAE", "00000000000191;6201501"])
            with self.assertRaisesRegex(ErroConsulta, "Layout"):
                buscar(Path(pasta) / "Estabelecimentos0.zip", {b"00000000000191"})

    def test_resultado_formata_codigos_e_remove_repetidos(self):
        tabela = {"6201501": "Desenvolvimento", "6204000": "Consultoria", "0111301": "Arroz"}
        campos = linha("00000000000191", "111301", "6204000,6201501,6204000,0").split(";")
        resultado = montar_resultado([c.strip('"') for c in campos], tabela)
        self.assertEqual(resultado["cnae_principal"], {"codigo": "01.11-3-01", "descricao": "Arroz"})
        self.assertEqual([a["codigo"] for a in resultado["cnaes_secundarios"]], ["62.04-0-00", "62.01-5-01"])
        with self.assertRaisesRegex(ErroConsulta, "ausente"):
            montar_resultado([c.strip('"') for c in linha("00000000000191", "9999999", "").split(";")], tabela)

    def test_lista_tamanhos_do_compartilhamento(self):
        xml = (b'<d:multistatus><d:response><d:href>/public.php/webdav/2026-09/</d:href></d:response>'
               b'<d:response><d:href>/public.php/webdav/2026-09/Cnaes.zip</d:href>'
               b'<d:getcontentlength>22078</d:getcontentlength></d:response></d:multistatus>')
        with patch("dados_abertos._abrir", return_value=io.BytesIO(xml)):
            self.assertEqual(_listar("2026-09/"), {"2026-09": None, "Cnaes.zip": 22078})


class Resposta(io.BytesIO):
    status = 206

    def __init__(self, dados: bytes, cair: bool):
        super().__init__(dados)
        self.cair = cair

    def read(self, n=-1):
        bloco = super().read(n)
        if not bloco and self.cair:
            raise ConnectionResetError("queda simulada")
        return bloco


class TestDownload(unittest.TestCase):
    def test_retoma_apos_queda_e_renomeia_ao_concluir(self):
        conteudo = b"0123456789"
        pedidos = []

        def abrir(caminho, metodo="GET", **cabecalhos):
            inicio = int(cabecalhos["Range"][6:-1])
            pedidos.append(inicio)
            return Resposta(conteudo[inicio:5] if inicio < 5 else conteudo[inicio:], cair=inicio < 5)

        with tempfile.TemporaryDirectory() as pasta, redirect_stderr(io.StringIO()):
            destino = Path(pasta) / "Estabelecimentos0.zip"
            with patch("dados_abertos._abrir", side_effect=abrir), patch("dados_abertos.time.sleep"):
                baixar("2026-09", destino.name, len(conteudo), destino)
            self.assertEqual(destino.read_bytes(), conteudo)
            self.assertFalse(destino.with_suffix(".parcial").exists())
        self.assertEqual(pedidos, [0, 5])


class TestImportar(unittest.TestCase):
    def test_preenche_excel_preserva_cache_recente_e_marca_ausentes(self):
        with tempfile.TemporaryDirectory() as pasta, redirect_stdout(io.StringIO()), redirect_stderr(io.StringIO()):
            raiz = Path(pasta)
            zips = raiz / "2026-09"
            zips.mkdir()
            criar_zips(zips, [linha("00000000000191", "6201501", "6204000"),
                              linha("12ABC34501DE35", "6204000", "")])
            entrada = raiz / "lista.txt"
            entrada.write_text("00.000.000/0001-91\n12ABC34501DE35\n11222333000181\n00000000000191\n")
            saida = raiz / "output.xlsx"
            anterior = {"cnae_principal": {"codigo": "62.01-5-01", "descricao": "Da BrasilAPI"},
                        "cnaes_secundarios": []}
            cache = sqlite3.connect(saida.with_suffix(".cache.sqlite3"))
            cache.execute("CREATE TABLE cnaes (cnpj TEXT PRIMARY KEY, consultado_em REAL, resultado TEXT)")
            cache.execute("INSERT INTO cnaes VALUES (?, ?, ?)", ("00000000000191", time.time(), json.dumps(anterior)))
            cache.commit()
            cache.close()
            self.assertEqual(importar(zips, entrada, saida), 1)  # Um CNPJ ausente segue pendente.
            livro = load_workbook(saida)
            linhas = [[c.value for c in row] for row in livro.active.iter_rows(min_row=2)]
            livro.close()
        self.assertEqual(linhas, [
            ["00.000.000/0001-91", "62.01-5-01 - Da BrasilAPI", "Não informada"],
            ["12.ABC.345/01DE-35", "62.04-0-00 - Consultoria em tecnologia da informação", "Não informada"],
            ["11.222.333/0001-81", "PENDENTE", "Não consultado"],
            ["00.000.000/0001-91", "62.01-5-01 - Da BrasilAPI", "Não informada"],
        ])


if __name__ == "__main__":
    unittest.main()
