from contextlib import redirect_stdout, redirect_stderr
from io import StringIO
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from consultar import ErroConsulta
from lote import CABECALHO, LimiteAPI, extrair_api, ler_cnpjs, montar_linhas, processar, salvar_excel

CNPJ = "00000000000191"
DADOS = {"cnpj": CNPJ, "cnae_fiscal": 6201501,
         "cnae_fiscal_descricao": "Desenvolvimento sob encomenda",
         "cnaes_secundarios": [{"codigo": 6204000, "descricao": "Consultoria"}],
         "razao_social": "DADO QUE NÃO DEVE SER EXPORTADO"}


class TestApi(unittest.TestCase):
    def test_extrai_apenas_os_cnaes(self):
        resultado = extrair_api(DADOS, CNPJ)
        self.assertEqual(set(resultado), {"cnae_principal", "cnaes_secundarios"})
        self.assertEqual(resultado["cnae_principal"]["codigo"], "62.01-5-01")
        self.assertNotIn("razao_social", str(resultado))

    def test_codigo_com_zero_inicial_e_sem_secundarios(self):
        dados = dict(DADOS, cnae_fiscal=111301, cnaes_secundarios=[])
        self.assertEqual(extrair_api(dados, CNPJ)["cnae_principal"]["codigo"], "01.11-3-01")
        self.assertEqual(extrair_api(dados, CNPJ)["cnaes_secundarios"], [])

    def test_secundarios_nao_informados_sentinela(self):
        dados = dict(DADOS, cnaes_secundarios=[{"codigo": 0, "descricao": ""}])
        self.assertEqual(extrair_api(dados, CNPJ)["cnaes_secundarios"], [])

    def test_resposta_incompleta_e_outro_cnpj_sao_erros(self):
        for dados in ({}, dict(DADOS, cnpj="12ABC34501DE35"),
                      dict(DADOS, cnaes_secundarios=None), dict(DADOS, cnae_fiscal_descricao="")):
            with self.subTest(dados=dados), self.assertRaises(ErroConsulta):
                extrair_api(dados, CNPJ)


class TestLote(unittest.TestCase):
    def test_csv_mantem_ordem_duplicatas_e_invalidos(self):
        with tempfile.TemporaryDirectory() as pasta:
            raiz = Path(pasta)
            (raiz / "lista.csv").write_text("Nome;CNPJ\nA;00.000.000/0001-91\nB;INVALIDO\nC;00000000000191\n")
            self.assertEqual(ler_cnpjs(raiz, raiz / "output.xlsx"),
                             ["00.000.000/0001-91", "INVALIDO", CNPJ])

    def test_xlsx_restaura_zeros_de_cnpj_numerico(self):
        with tempfile.TemporaryDirectory() as pasta:
            raiz = Path(pasta)
            livro = Workbook()
            livro.active.append(["Empresa", "CNPJ"])
            livro.active.append(["Exemplo", 191])
            livro.save(raiz / "lista.xlsx")
            self.assertEqual(ler_cnpjs(raiz, raiz / "output.xlsx"), [CNPJ])

    def test_txt_sem_cabecalho(self):
        with tempfile.TemporaryDirectory() as pasta:
            raiz = Path(pasta)
            (raiz / "lista.txt").write_text("00.000.000/0001-91\n12.ABC.345/01DE-35\n")
            self.assertEqual(len(ler_cnpjs(raiz, raiz / "output.xlsx")), 2)

    def test_excel_tem_tres_colunas_cnpj_texto_e_formula_inofensiva(self):
        with tempfile.TemporaryDirectory() as pasta:
            caminho = Path(pasta) / "output.xlsx"
            salvar_excel(caminho, [[CNPJ, "=1+1", "Não informada"]])
            livro = load_workbook(caminho)
            self.assertEqual([c.value for c in livro.active[1]], CABECALHO)
            self.assertEqual(livro.active["A2"].data_type, "s")
            self.assertEqual(livro.active["B2"].data_type, "s")
            self.assertEqual(livro.active.freeze_panes, "B2")
            livro.close()

    def test_consulta_duplicada_uma_vez_e_retoma_do_cache(self):
        with tempfile.TemporaryDirectory() as pasta, redirect_stdout(StringIO()):
            raiz = Path(pasta)
            entrada = raiz / "lista.txt"
            entrada.write_text("00.000.000/0001-91\n00000000000191\n")
            with patch("lote.consultar_api", return_value=extrair_api(DADOS, CNPJ)) as api:
                self.assertEqual(processar(entrada, raiz / "output.xlsx"), 0)
                self.assertEqual(processar(entrada, raiz / "output.xlsx"), 0)
                api.assert_called_once_with(CNPJ)
            livro = load_workbook(raiz / "output.xlsx")
            self.assertEqual(livro.active.max_row, 3)
            livro.close()

    def test_bloqueio_da_api_preserva_pendentes_e_interrompe(self):
        with tempfile.TemporaryDirectory() as pasta, redirect_stdout(StringIO()), redirect_stderr(StringIO()):
            raiz = Path(pasta)
            entrada = raiz / "lista.txt"
            entrada.write_text("00000000000191\n12ABC34501DE35\n")
            with patch("lote.consultar_api", side_effect=LimiteAPI("HTTP 429")) as api:
                self.assertEqual(processar(entrada, raiz / "output.xlsx"), 1)
                api.assert_called_once()
            livro = load_workbook(raiz / "output.xlsx")
            self.assertIn("ERRO:", livro.active["B2"].value)
            self.assertEqual(livro.active["B3"].value, "PENDENTE")
            livro.close()

    def test_recusa_sobrescrever_planilha_estranha(self):
        with tempfile.TemporaryDirectory() as pasta:
            raiz = Path(pasta)
            entrada = raiz / "lista.txt"
            entrada.write_text(CNPJ)
            livro = Workbook()
            livro.active.append(["OUTROS DADOS"])
            livro.save(raiz / "output.xlsx")
            with self.assertRaisesRegex(ErroConsulta, "não é deste projeto"):
                processar(entrada, raiz / "output.xlsx")


if __name__ == "__main__":
    unittest.main()
